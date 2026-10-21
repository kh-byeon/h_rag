#!/usr/bin/env python3
"""Workbook writer for scripts/analyze_q1000.py. Not meant to be run directly."""
from __future__ import annotations

import datetime as _dt

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="595959")
TITLE = Font(name=FONT, size=13, bold=True, color="1F3864")
HIT = PatternFill("solid", fgColor="E2EFDA")   # 95% CI excludes zero
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(bottom=THIN)

# Column display names and number formats.
FMT = {
    "n": "0", "l2": "0.0000", "mmr": "0.0000", "delta": "+0.0000;-0.0000;0.0000",
    "rel_change": "+0.0%;-0.0%;0.0%", "lo": "+0.0000;-0.0000;0.0000",
    "hi": "+0.0000;-0.0000;0.0000", "p_t": "0.00E+00", "p_wilcoxon": "0.00E+00",
    "mcnemar_p": "0.00E+00", "tost_pm0_02": "0.0000", "effect": "+0.0000;-0.0000;0.0000",
    "sd_per_query": "0.000", "ci_halfwidth_at_n1000": "0.0000",
    "n_for_80pct_power": "#,##0", "pct_at_ceiling": "0.0", "cutoff": "0.00",
    "l2_f1": "0.0000", "mmr_f1": "0.0000", "l2_em": "0.0", "mmr_em": "0.0",
    "l2_rec": "0.000", "mmr_rec": "0.000",
}
HEAD = {
    "l2": "L2", "mmr": "MMR", "delta": "Δ (MMR−L2)", "rel_change": "rel. change",
    "lo": "95% CI lo", "hi": "95% CI hi", "p_t": "p (paired t)",
    "p_wilcoxon": "p (Wilcoxon)", "sig95": "CI excl. 0?", "bonferroni": f"Bonferroni",
    "mcnemar_l2_only": "McNemar: L2-only", "mcnemar_mmr_only": "McNemar: MMR-only",
    "mcnemar_p": "McNemar p", "tost_pm0_02": "TOST p (±0.02)",
    "n_for_80pct_power": "n for 80% power", "sd_per_query": "SD per query",
    "ci_halfwidth_at_n1000": "CI half-width @n=1000", "pct_at_ceiling": "% at ceiling",
}


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _write_table(ws, df, start_row, *, rel_change=True):
    """Write df; delta and rel_change become live formulas on the L2/MMR cells."""
    cols = list(df.columns)
    if rel_change and {"l2", "mmr", "delta"} <= set(cols):
        cols.insert(cols.index("delta") + 1, "rel_change")
    for j, name in enumerate(cols, start=1):
        ws.cell(row=start_row, column=j, value=HEAD.get(name, name))
    _style_header(ws, start_row, len(cols))

    col_of = {name: get_column_letter(j) for j, name in enumerate(cols, start=1)}
    for i, (_, rec) in enumerate(df.iterrows()):
        r = start_row + 1 + i
        for j, name in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=j)
            cell.font = BODY
            cell.border = BOX
            if name == "delta" and "l2" in col_of and "mmr" in col_of:
                # live formula, so the sheet recomputes if a mean is corrected
                cell.value = f"=IF(OR({col_of['l2']}{r}=\"\",{col_of['mmr']}{r}=\"\"),\"\",{col_of['mmr']}{r}-{col_of['l2']}{r})"
            elif name == "rel_change":
                cell.value = f"=IF(OR({col_of['l2']}{r}=\"\",{col_of['l2']}{r}=0),\"\",{col_of['delta']}{r}/{col_of['l2']}{r})"
            else:
                v = rec.get(name, "")
                if isinstance(v, (np.floating, np.integer)):
                    v = float(v) if isinstance(v, np.floating) else int(v)
                if isinstance(v, float) and np.isnan(v):
                    v = ""
                cell.value = v
            if name in FMT:
                cell.number_format = FMT[name]
            if name in ("sig95", "bonferroni", "quantization", "metric", "group",
                        "contrast", "split", "quantity"):
                cell.alignment = Alignment(horizontal="left")
        if str(rec.get("sig95", "")) == "yes" or str(rec.get("bonferroni", "")) == "pass":
            for j in range(1, len(cols) + 1):
                ws.cell(row=r, column=j).fill = HIT
    return start_row + len(df) + 1, cols


def _autosize(ws, cols, widths=None):
    for j, name in enumerate(cols, start=1):
        w = (widths or {}).get(name)
        if w is None:
            w = max(11, min(34, len(str(HEAD.get(name, name))) + 4))
        ws.column_dimensions[get_column_letter(j)].width = w


def _sheet(wb, title, heading, blurb, df, *, rel_change=True, widths=None):
    ws = wb.create_sheet(title)
    ws.cell(row=1, column=1, value=heading).font = TITLE
    r = 2
    for line in blurb:
        ws.cell(row=r, column=1, value=line).font = NOTE
        r += 1
    end, cols = _write_table(ws, df, r + 1, rel_change=rel_change)
    _autosize(ws, cols, widths)
    ws.cell(row=end + 1, column=1,
            value="Green = the 95% bootstrap CI excludes zero (or the Bonferroni column passes).").font = NOTE
    return ws


def write_workbook(out, path, *, run_id, seed, boots, n_tests):
    wb = Workbook()
    wb.remove(wb.active)

    # ------------------------------------------------------------- README --
    ws = wb.create_sheet("README")
    ws.cell(row=1, column=1, value="H_RAG v3.0 — on-the-fly MMR, n=1000 distractor run").font = Font(
        name=FONT, size=15, bold=True, color="1F3864")
    lines = [
        ("", ""),
        ("What this is", "Every published metric for the n=1000 HotpotQA distractor run, L2 baseline vs on-the-fly MMR."),
        ("Run ID", run_id),
        ("Configuration", "distractor scope · λ=0.7 · τ=−1.0 · γ=0.7 · top_k=5 · one combo · 1000 queries"),
        ("Device path", "FeFET centroid (fefet_ref_fit_smu.json + linear_sum, V_ML=1.4) + IGZO capacitor MMR"),
        ("Generator", "Meta-Llama-3-8B-Instruct, greedy (do_sample=False), max_new_tokens=512, at fp16 / int8 / int4"),
        ("", ""),
        ("HOW TO REGENERATE", ""),
        ("  1. produce the data", "bash scripts/run_q1000_distractor.sh        (idempotent; skips whatever is already done)"),
        ("  2. rebuild this file", "python scripts/analyze_q1000.py --rebuild   (writes figures/q1000_metrics.xlsx)"),
        ("  fast path", "python scripts/analyze_q1000.py             (reuses figures/perq_em_f1_q1000.csv)"),
        ("  written up in", "docs/ANALYSIS_q1000.md  ·  findings in docs/RESULTS_q1000.md"),
        ("", ""),
        ("STATISTICS — read before quoting anything", ""),
        ("  unit of analysis", "the QUERY. Each metric is reduced to one value per query_idx, then bootstrapped over the 1000 query means."),
        ("", "Testing over rows (6,000 of them) is pseudo-replication and has produced absurd p-values in this project before."),
        ("  intervals", f"paired bootstrap, {boots:,} resamples, numpy default_rng(seed={seed}). Reproducible to the digit."),
        ("  p-values", "two-sided paired t-test and Wilcoxon signed-rank, reported side by side; they agree throughout."),
        ("  Bonferroni column", f"'pass' = both p < 0.05/{n_tests}, the family size used across this analysis."),
        ("  green shading", "the 95% CI excludes zero, or the Bonferroni column passes."),
        ("", ""),
        ("PROVENANCE OF EACH NUMBER", ""),
        ("  L2 / MMR means", "computed from per-query data; Δ and rel. change are live formulas on those two cells."),
        ("  CI, p-values, TOST", "computed by scripts/analyze_q1000.py and written as values — they cannot be expressed as sheet formulas."),
        ("  ILD / RR / breadth / ROUGE-L", "copied from sweep_results_full_grid.csv, which stores CORPUS AGGREGATES ONLY."),
        ("", "Those four rows therefore have no CI and are not testable without a per-query re-export."),
        ("  EM / F1 identity check", "the per-query table reproduces figures/em_f1_q1000.csv to 1e-9 on all six arm×quant means;"),
        ("", "analyze_q1000.py asserts this on every run and aborts if it fails."),
        ("", ""),
        ("KNOWN LIMITS", ""),
        ("  single combo", "one hyperparameter setting; says nothing about robustness across the grid."),
        ("  single model / dataset", "one generator, one dataset, one seed (greedy decoding, so seeds do not apply)."),
        ("  FeFET is dead compute here", "under distractor scope the centroid search does not affect candidate selection (HANDOFF §3)."),
        ("", "Latency and energy are measurable; accuracy is not attributable to the device."),
        ("  DiD is under-powered", "the quantization-robustness interaction is not resolved at n=1000 — see 'Key contrasts' and 'Power'."),
        ("", ""),
        ("Generated", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (k, v) in enumerate(lines, start=2):
        a = ws.cell(row=i, column=1, value=k)
        a.font = BOLD if (k and not k.startswith("  ")) else BODY
        b = ws.cell(row=i, column=2, value=v)
        b.font = BODY
        b.alignment = Alignment(wrap_text=False)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 118

    # ------------------------------------------------------------- sheets --
    _sheet(wb, "Headline", "MMR vs L2, at each precision",
           ["The primary comparison. Each row is a paired test over 1000 queries at one precision.",
            "n_no_info counts 'the passage does not mention X' statements per answer — LOWER is better, so a negative Δ is an improvement.",
            "answer_relevance is retained only to document that it is uninformative here (|Δ| ≤ 0.007)."],
           out["headline"], widths={"metric": 26, "quantization": 14})

    _sheet(wb, "Retrieval", "Retrieval quality (LLM-independent)",
           ["Measured once, before generation — these numbers do not depend on the generator or its precision.",
            "'full support' = both gold paragraphs are inside the top-5 handed to the LLM.",
            "The last four rows come from the sweep CSV, which stores aggregates only: no CI is available for them."],
           out["retrieval"], widths={"metric": 42})

    _sheet(wb, "Key contrasts", "Deployment contrasts and the quantization question",
           ["MMR@int4 vs L2@fp16 — the deployable comparison: does MMR at edge precision match the fp16 baseline?",
            "TOST is a two-one-sided-test for equivalence within ±0.02; the margin was chosen post hoc, so treat it as descriptive.",
            "DiD rows ask whether MMR degrades LESS than L2 from fp16 to int4. Both are NOT significant at n=1000."],
           out["contrasts"], rel_change=False, widths={"contrast": 46})

    _sheet(wb, "Mechanism", "Where the accuracy gain comes from",
           ["The F1 delta split by which arm retrieved both gold paragraphs. Retrieval is identical across precisions,",
            "so this grouping is independent of the quantization axis.",
            "RETRIEVAL-MATCHED is the key row: when both arms see the same evidence, the arms generate equally well."],
           out["mechanism"], widths={"group": 34, "quantization": 14})

    _sheet(wb, "Subgroups", "F1 delta by question type and difficulty",
           ["HotpotQA dev-distractor is uniformly 'hard', so the level split has a single group.",
            "Bridge vs comparison show the same effect; comparison is simply under-powered at n=193."],
           out["subgroups"], widths={"group": 16, "split": 10, "quantization": 14})

    _sheet(wb, "Faithfulness cutoff", "Sensitivity to the faithfulness cutoff",
           ["Re-binarises the CONTINUOUS mean-entailment score at each cutoff. This approximates --threshold,",
            "which acts per claim, so it is a sensitivity check and not a substitute for re-running the NLI scorer.",
            "The int4 effect sharpens as the cutoff rises while fp16 stays flat — the de-saturation is not what drives it."],
           out["faith_cutoff"], widths={"quantization": 14})

    _sheet(wb, "Power", "Observed variance and required sample size",
           ["n for 80% power assumes the observed point estimate is the true effect — it is a planning figure, not a guarantee.",
            "If the true effect is smaller than measured, the required n grows quadratically.",
            "The two DiD rows are why the quantization-robustness claim is not currently supported."],
           out["power"], rel_change=False, widths={"quantity": 30})

    ws = wb.create_sheet("PerQuery")
    ws.cell(row=1, column=1, value="Per-query source data (3 precisions × 1000 queries)").font = TITLE
    ws.cell(row=2, column=1, value="Every table in this workbook is computed from these rows. "
                                   "l2_full / mmr_full = both gold paragraphs in the top-5.").font = NOTE
    df = out["perquery"].copy()
    for j, name in enumerate(df.columns, start=1):
        ws.cell(row=4, column=j, value=name)
    _style_header(ws, 4, len(df.columns))
    for i, rec in enumerate(df.itertuples(index=False)):
        for j, v in enumerate(rec, start=1):
            cell = ws.cell(row=5 + i, column=j)
            if isinstance(v, (np.floating, np.integer)):
                v = float(v) if isinstance(v, np.floating) else int(v)
            if isinstance(v, float) and np.isnan(v):
                v = ""
            cell.value = v
            cell.font = BODY
            if df.columns[j - 1] in FMT:
                cell.number_format = FMT[df.columns[j - 1]]
    _autosize(ws, list(df.columns), {"query": 60})
    ws.auto_filter.ref = f"A4:{get_column_letter(len(df.columns))}{4 + len(df)}"

    wb.save(path)
    return path
